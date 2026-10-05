import os
import tkinter as tk
from tkinter import ttk, filedialog
from typing import Any, Optional

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

import matplotlib as mpl
import matplotlib.dates as mdates
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg, NavigationToolbar2Tk
from matplotlib.figure import Figure

from Buttons import MainButtons, TestButtons, SmallButtons
from Messages import Messages
from Settings import (
    show_info, FONTS, COLORS, GAPS, SPLITTER_MINSIZES,
    PLOT, TABLE, ENTRY_WIDTH, IMAX_BASE_RANGE
)

mpl.rcParams['font.family'] = 'Times New Roman'



class Imax_processing:
    """Find current values corresponding to a requested reactor heating rate.

    The user selects the heating interval on the temperature plot, enters the
    requested heating rate (degC/h) and a divisor. For every current column the
    program finds the point of the interval whose heating rate is the closest
    to the requested one and outputs the current at that point divided by the
    divisor (Imax). If MC (heat capacity, MWt/(°C/h)) was given in the module,
    Pmax = Imax * MC is also calculated for every current column.
    The window (base, s) for the heating rate is also set in the module.
    """

    def __init__(self, root: tk.Widget, main_app: Any) -> None:
        self.root = root
        self.main_app = main_app
        self.Imax_processing_frame: Optional[tk.Frame] = None
        self.menu_bar: Optional[tk.Menu] = None

        self.Time = getattr(main_app, 'Time', None)
        self.Temperature = getattr(main_app, 'Temperature', None)
        self.Current = getattr(main_app, 'Current', None)
        self.Current_columns = list(getattr(main_app, 'Current_columns', []))
        # Window of the heating rate, s (default 60) and optional heat capacity
        self.base: int = int(getattr(main_app, 'base', 60) or 60)
        self.MC: Optional[float] = getattr(main_app, 'MC', None)

        self.selection_left: Optional[int] = None
        self.selection_right: Optional[int] = None
        self.selection_artists = []
        # One entry per table row: the search parameters and the found row.
        # They are kept so that the rows can be recalculated when MC / base
        # are edited.
        self.steps: list[dict[str, Any]] = []
        self.point_marker: Optional[Any] = None
        self.active_entry: Optional[ttk.Entry] = None
        self._zoom_factor = 1.25

    @property
    def results(self) -> list[dict[str, Any]]:
        """Rows currently shown in the table."""
        return [step['row'] for step in self.steps]

    def create_Imax_processing_window(self) -> None:
        root_window = self.root.winfo_toplevel()
        root_window.config(menu=None)
        self.Imax_processing_frame = tk.Frame(
            self.root, bg=COLORS['BACKGROUND_COLOR']
        )
        self.Imax_processing_frame.place(x=0, y=0, relwidth=1, relheight=1)
        root_window.title('Imax')

        self.Imax_processing_frame.grid_rowconfigure(1, weight=1)
        self.Imax_processing_frame.grid_columnconfigure(0, weight=4)
        self.Imax_processing_frame.grid_columnconfigure(1, weight=6)

        self.create_menu()
        self.create_splitter_window()
        self.create_plot()
        self.create_table()
        self.create_controls()
        self.create_buttons()

    def create_menu(self) -> None:
        root_window = self.root.winfo_toplevel()
        self.menu_bar = tk.Menu(root_window)
        root_window.config(menu=self.menu_bar)

        info = tk.Menu(self.menu_bar, tearoff=1)
        self.menu_bar.add_cascade(label='INFO', menu=info)
        info.add_command(
            label='INFO', font=FONTS['DATA_FONT'],
            command=lambda: show_info(self.root, 'Imax/info_processing.txt')
        )

        computed = tk.Menu(self.menu_bar, tearoff=1)
        self.menu_bar.add_cascade(label='Computed parameters', menu=computed)
        computed.add_command(
            label='Show computed parameters', font=FONTS['DATA_FONT'],
            command=self.show_computed_parameters
        )
        computed.add_command(
            label='Edit computed parameters', font=FONTS['DATA_FONT'],
            command=self.edit_computed_parameters
        )

        save = tk.Menu(self.menu_bar, tearoff=1)
        self.menu_bar.add_cascade(label='Save', menu=save)
        save.add_command(
            label='Save to Excel', font=FONTS['DATA_FONT'], command=self.save_to_excel
        )
        save.add_command(
            label="Save in '.txt'", font=FONTS['DATA_FONT'], command=self.save_to_txt
        )

    def create_splitter_window(self) -> None:
        self.splitter_window = tk.PanedWindow(
            self.Imax_processing_frame, orient=tk.HORIZONTAL,
            sashrelief=tk.RAISED, sashwidth=5,
            bg=COLORS['BACKGROUND_COLOR']
        )
        self.splitter_window.grid(
            row=1, column=0, columnspan=2,
            padx=GAPS['GAPS_X']['PAD_X_10'],
            pady=GAPS['GAPS_Y']['PAD_Y_10_5'], sticky='nsew'
        )

        self.table_frame = tk.Frame(
            self.splitter_window, bg=COLORS['PARAMETERS_BACKGROUND_COLOR']
        )
        self.plot_frame = tk.Frame(
            self.splitter_window, bg=COLORS['BACKGROUND_COLOR']
        )
        self.splitter_window.add(
            self.table_frame, minsize=SPLITTER_MINSIZES['LEFT'], stretch='always'
        )
        self.splitter_window.add(
            self.plot_frame, minsize=SPLITTER_MINSIZES['RIGHT'], stretch='always'
        )

    @staticmethod
    def _to_datetimes(values: pd.Series) -> pd.DatetimeIndex:
        """Convert the time column to real clock time.

        Numeric values are treated as Excel serial dates (days since
        1899-12-30), exactly as in the DRDH module. Datetime columns and
        date/time strings are converted as they are.
        """
        if pd.api.types.is_datetime64_any_dtype(values):
            dt = pd.to_datetime(values, errors='coerce')
        else:
            numeric = pd.to_numeric(values, errors='coerce')
            if numeric.notna().all():
                dt = (
                    pd.Timestamp('1899-12-30')
                    + pd.to_timedelta(numeric.to_numpy(float), unit='D')
                )
            else:
                dt = pd.to_datetime(values, errors='coerce')

        dt = pd.DatetimeIndex(dt)
        if dt.isna().any():
            raise ValueError(
                'Time column contains values that cannot be interpreted as time.'
            )
        if dt.tz is not None:
            dt = dt.tz_localize(None)

        # Float days carry micro-second noise; the data are 1 Hz anyway.
        return dt.round('ms')

    @staticmethod
    def _format_time(x: float) -> str:
        """Matplotlib date number -> 'HH:MM:SS'."""
        return mdates.num2date(x).strftime('%H:%M:%S')

    def create_plot(self) -> None:
        # Real clock time is shown on the X axis, like in the other modules.
        times = self._to_datetimes(self.Time)
        self.times = times
        self.time_num = mdates.date2num(times.to_numpy())

        self.temperature_values = pd.to_numeric(
            self.Temperature, errors='coerce'
        ).to_numpy(float)

        if np.isnan(self.temperature_values).any():
            raise ValueError('Temperature contains invalid values.')
        if len(self.time_num) < 2:
            raise ValueError('At least two data points are required.')

        # Heating rate of every point, °C/h. Computed once: it is used both
        # for the optional curve on the plot and by Proceed.
        self.rate_values = self._heating_rate_per_hour(
            self.temperature_values, self.base
        )

        self.fig = Figure(figsize=(7, 5), constrained_layout=True)
        self.ax = self.fig.add_subplot(111)
        (self.temperature_line,) = self.ax.plot(
            self.times, self.temperature_values,
            linewidth=1.3, label='Temperature'
        )
        self.ax.set_xlabel('Time', fontsize=PLOT['PLOT_LABEL_SIZE'])
        self.ax.set_ylabel('Temperature, °C', fontsize=PLOT['PLOT_LABEL_SIZE'])
        self.ax.grid(True, alpha=0.35)

        # Temperature rise speed: second Y axis, hidden until the checkbox
        # "Display temperature rise speed" is ticked.
        self.rate_color = 'tab:red'
        self.ax2 = self.ax.twinx()
        (self.rate_line,) = self.ax2.plot(
            self.times, self.rate_values,
            color=self.rate_color, linewidth=1.0, alpha=0.8,
            label='Temperature rise speed'
        )
        self.ax2.set_ylabel(
            'Temperature rise speed, °C/h',
            fontsize=PLOT['PLOT_LABEL_SIZE']
        )
        # Requested heating rate (the 'Heating rate' entry), dashed. Its
        # position is set by update_target_line once the entry exists.
        self.target_line = self.ax2.axhline(
            0.0, color=self.rate_color, linestyle='--', linewidth=1.2,
            label='Target speed'
        )
        self.ax2.set_visible(False)
        self.ax2.set_in_layout(False)

        # Currents: third Y axis, shifted outwards so that it does not
        # overlap the speed axis. Hidden until "Display currents" is ticked.
        # The raw values from the file are drawn, not divided ones.
        palette = [
            'tab:green', 'tab:purple', 'tab:brown', 'tab:pink',
            'tab:olive', 'tab:cyan', 'tab:gray'
        ]
        self.ax3 = self.ax.twinx()
        self.ax3.spines['right'].set_position(('outward', 65))
        for side in ('left', 'top', 'bottom'):
            self.ax3.spines[side].set_visible(False)

        self.current_lines = []
        for i, column in enumerate(self.Current_columns):
            values = pd.to_numeric(
                self.Current[column], errors='coerce'
            ).to_numpy(float)
            (line,) = self.ax3.plot(
                self.times, values,
                color=palette[i % len(palette)], linewidth=1.0, alpha=0.9,
                label=str(column)
            )
            self.current_lines.append(line)

        self.ax3.set_ylabel('Current', fontsize=PLOT['PLOT_LABEL_SIZE'])
        self.ax3.set_visible(False)
        self.ax3.set_in_layout(False)

        # Every axis that may receive mouse events
        self.plot_axes = (self.ax, self.ax2, self.ax3)

        # Keep the temperature axis (curve, legend, selection) on top
        self.ax.set_zorder(max(self.ax2.get_zorder(), self.ax3.get_zorder()) + 1)
        self.ax.patch.set_visible(False)
        self.update_legend()

        # The date is added only if the experiment lasts longer than a day.
        span_days = float(self.time_num.max() - self.time_num.min())
        fmt = '%d.%m %H:%M:%S' if span_days > 1 else '%H:%M:%S'
        self.ax.xaxis.set_major_formatter(mdates.DateFormatter(fmt))
        self.fig.autofmt_xdate()

        self.canvas = FigureCanvasTkAgg(self.fig, master=self.plot_frame)
        self.canvas_widget = self.canvas.get_tk_widget()
        self.canvas_widget.pack(fill=tk.BOTH, expand=True)
        self.toolbar = NavigationToolbar2Tk(self.canvas, self.plot_frame)
        self.toolbar.update()
        self.toolbar.pack(side=tk.TOP, fill=tk.X)

        self.canvas.mpl_connect('button_press_event', self.on_plot_click)
        self.canvas.mpl_connect('scroll_event', self.on_scroll)

    def update_legend(self) -> None:
        """Legend with the temperature and, if shown, the speed curve."""
        lines = [self.temperature_line]
        if self.ax2.get_visible():
            lines.append(self.rate_line)
            lines.append(self.target_line)
        if self.ax3.get_visible():
            lines.extend(self.current_lines)

        legend = self.ax.get_legend()
        if legend is not None:
            legend.remove()
        self.ax.legend(lines, [line.get_label() for line in lines], loc='best')

    def toggle_rate_display(self) -> None:
        """Show or hide the temperature rise speed on the plot."""
        show = self.rate_var.get() == 1
        self.ax2.set_visible(show)
        self.ax2.set_in_layout(show)
        self.update_legend()
        self.canvas.draw_idle()

    def currents_shown(self) -> bool:
        """True while the 'Display currents' checkbox is ticked."""
        return hasattr(self, 'currents_var') and self.currents_var.get() == 1

    def toggle_currents_display(self) -> None:
        """Show or hide the current curves and their (third) Y axis."""
        show = self.currents_shown()
        self.ax3.set_visible(show)
        self.ax3.set_in_layout(show)
        self.update_legend()
        self.canvas.draw_idle()

    def refresh_point_marker(self) -> None:
        """Dashed black line at the point of the last table row.

        Only one line exists: it is replaced by every new search and moved
        back to the previous row when a step is removed.
        """
        if self.point_marker is not None:
            try:
                self.point_marker.remove()
            except ValueError:
                pass
            self.point_marker = None
        if self.steps:
            x = self.time_num[self.steps[-1]['source_index']]
            self.point_marker = self.ax.axvline(
                x, color='black', linestyle='--', linewidth=1.2, zorder=5
            )
        self.canvas.draw_idle()

    def update_target_line(self) -> None:
        """Move the dashed 'target speed' line to the value of the entry."""
        try:
            value = float(self.speed_entry.get().strip().replace(',', '.'))
        except ValueError:
            return
        self.target_line.set_ydata([value, value])
        self.ax2.relim()
        self.ax2.autoscale_view(scalex=False)
        self.canvas.draw_idle()

    def create_table(self) -> None:
        # Pmax columns exist only if MC was entered
        self.Pmax_columns = (
            [f'Pmax {c}' for c in self.Current_columns] if self.MC is not None else []
        )
        columns = ['speed'] + self.Current_columns + self.Pmax_columns
        self.columns = columns
        style = ttk.Style(self.root)
        style.configure('Imax.Treeview', font=FONTS['DATA_FONT'])
        style.configure('Imax.Treeview.Heading', font=FONTS['HEADING_FONT'])

        self.tree = ttk.Treeview(
            self.table_frame, columns=columns, show='headings',
            style='Imax.Treeview', height=TABLE['CELL_HEIGHT']
        )
        for column in columns:
            if column == 'speed':
                heading = 'Temperature rise speed, °C/h'
            elif column in self.Current_columns:
                heading = f'Imax {column}'
            else:
                heading = column
            self.tree.heading(column, text=heading)
            self.tree.column(column, width=TABLE['CELL_WIDTH'], anchor='center')

    def create_controls(self) -> None:
        controls = tk.Frame(
            self.Imax_processing_frame, bg=COLORS['BACKGROUND_COLOR']
        )
        controls.grid(
            row=2, column=0, columnspan=2,
            padx=GAPS['GAPS_X']['PAD_X_10'], pady=GAPS['GAPS_Y']['PAD_Y_5'],
            sticky='ew'
        )
        controls.grid_columnconfigure(6, weight=1)

        tk.Label(
            controls, text='Heating rate, °C/h:', font=FONTS['TEXT_FONT'],
            bg=COLORS['BACKGROUND_COLOR']
        ).grid(row=0, column=0, padx=5)
        self.speed_entry = ttk.Entry(controls, width=ENTRY_WIDTH, font=FONTS['DATA_FONT'])
        self.speed_entry.grid(row=0, column=1, padx=5)
        self.speed_entry.insert(0, '10')
        self.speed_entry.bind('<KeyRelease>', lambda e: self.update_target_line())
        self.update_target_line()

        tk.Label(
            controls, text='Divide current by:', font=FONTS['TEXT_FONT'],
            bg=COLORS['BACKGROUND_COLOR']
        ).grid(row=0, column=2, padx=5)
        self.divide_entry = ttk.Entry(controls, width=ENTRY_WIDTH, font=FONTS['DATA_FONT'])
        self.divide_entry.grid(row=0, column=3, padx=5)
        self.divide_entry.insert(0, '60')

        self.selection_label = tk.Label(
            controls, text='Select heating interval: left click — first point, right click — last point',
            font=FONTS['INFO_FONT'], bg=COLORS['BACKGROUND_COLOR']
        )
        self.selection_label.grid(row=0, column=4, columnspan=3, padx=10, sticky='w')

        self.rate_var = tk.IntVar(value=0)
        self.rate_checkbox = tk.Checkbutton(
            controls, text='Display temperature rise speed',
            variable=self.rate_var, command=self.toggle_rate_display,
            bg=COLORS['BACKGROUND_COLOR'], font=FONTS['TEXT_FONT']
        )
        self.rate_checkbox.grid(
            row=1, column=0, columnspan=4, padx=5, sticky='w'
        )

        self.currents_var = tk.IntVar(value=0)
        self.currents_checkbox = tk.Checkbutton(
            controls, text='Display currents',
            variable=self.currents_var, command=self.toggle_currents_display,
            bg=COLORS['BACKGROUND_COLOR'], font=FONTS['TEXT_FONT']
        )
        self.currents_checkbox.grid(
            row=2, column=0, columnspan=4, padx=5, sticky='w'
        )

    def create_buttons(self) -> None:
        def place(button, row: int, column: int) -> None:
            button.grid(
                row=row, column=column,
                padx=GAPS['GAPS_X']['PAD_X_40_20'],
                pady=GAPS['GAPS_Y']['PAD_Y_10'], sticky='w'
            )

        self.Proceed_button = TestButtons(
            self.Imax_processing_frame, text='Proceed', command=self.proceed
        )
        place(self.Proceed_button, 3, 0)
        self.remove_button = TestButtons(
            self.Imax_processing_frame, text='Remove previous step',
            command=self.remove_previous_step
        )
        place(self.remove_button, 3, 1)
        self.save_button = TestButtons(
            self.Imax_processing_frame, text='Save', command=self.choose_save_format
        )
        place(self.save_button, 4, 0)
        self.INFO_button = MainButtons(
            self.Imax_processing_frame, text='User guide',
            command=lambda: show_info(self.root, 'Imax/info_processing.txt')
        )
        place(self.INFO_button, 4, 1)
        self.BACK_button = MainButtons(
            self.Imax_processing_frame, text='<< BACK',
            command=lambda: self.back(self.Imax_processing_frame)
        )
        place(self.BACK_button, 5, 0)

    def _nearest_index(self, x: float) -> int:
        return int(np.argmin(np.abs(self.time_num - x)))

    def on_plot_click(self, event) -> None:
        # Do not set points while zoom / pan of the toolbar is active.
        if self.toolbar.mode != '':
            return
        if event.inaxes not in self.plot_axes or event.xdata is None:
            return
        idx = self._nearest_index(float(event.xdata))
        if event.button == 1:
            self.selection_left = idx
        elif event.button == 3:
            self.selection_right = idx
        else:
            return

        if self.selection_left is not None and self.selection_right is not None:
            if self.selection_left > self.selection_right:
                self.selection_left, self.selection_right = (
                    self.selection_right, self.selection_left
                )
            self.selection_label.config(
                text=(
                    f'Selected: {self._format_time(self.time_num[self.selection_left])}'
                    f' – {self._format_time(self.time_num[self.selection_right])}'
                )
            )
        else:
            side = 'first point' if event.button == 1 else 'last point'
            self.selection_label.config(text=f'{side.capitalize()} point selected.')
        self.redraw_selection()

    def redraw_selection(self) -> None:
        for artist in self.selection_artists:
            try:
                artist.remove()
            except ValueError:
                pass
        self.selection_artists = []
        if self.selection_left is not None:
            self.selection_artists.append(
                self.ax.axvline(
                    self.time_num[self.selection_left],
                    linestyle='--', linewidth=1.2
                )
            )
        if self.selection_right is not None:
            self.selection_artists.append(
                self.ax.axvline(
                    self.time_num[self.selection_right],
                    linestyle='--', linewidth=1.2
                )
            )
        if self.selection_left is not None and self.selection_right is not None:
            self.selection_artists.append(
                self.ax.axvspan(
                    self.time_num[self.selection_left],
                    self.time_num[self.selection_right], alpha=0.15
                )
            )
        self.canvas.draw_idle()

    def _read_parameters(self) -> tuple[float, float]:
        try:
            speed = float(self.speed_entry.get().strip().replace(',', '.'))
            divide = float(self.divide_entry.get().strip().replace(',', '.'))
        except ValueError as exc:
            raise ValueError('Heating rate and divisor must be numeric.') from exc
        if speed < 0:
            raise ValueError('Heating rate must be non-negative.')
        if divide <= 0:
            raise ValueError('Divisor must be greater than zero.')
        return speed, divide

    @staticmethod
    def _heating_rate_per_hour(
        temperature: np.ndarray, window: int = 60
    ) -> np.ndarray:
        """Return the heating rate in °C/h for every point.

        Measurements are recorded once per second, so a `window`-point span
        is `window` seconds. For an interior point i the rate is centered:
            (T[i + window/2] - T[i - window/2]) / window s * 3600
        e.g. window = 60: (T[i + 30] - T[i - 30]) / 60 * 3600;
             window = 30: (T[i + 15] - T[i - 15]) / 30 * 3600.
        Near the ends of the series the same-length span is shifted inside
        the data (one-sided difference). If the series is shorter than the
        window, the whole series is used.
        """
        n = len(temperature)
        if n < 2:
            raise ValueError('At least two data points are required.')

        span = min(window, n - 1)
        lo = np.arange(n) - span // 2
        hi = lo + span

        # Shift the span inside [0, n - 1]
        shift = np.where(lo < 0, -lo, 0)
        lo, hi = lo + shift, hi + shift
        over = np.where(hi > n - 1, hi - (n - 1), 0)
        lo, hi = lo - over, hi - over

        return (temperature[hi] - temperature[lo]) / span * 3600.0

    def _compute_step(
        self, left: int, right: int, target: float, divide: float,
        rate_values: np.ndarray, mc: Optional[float]
    ) -> tuple[dict[str, float], int]:
        """Find the row for the interval [left, right]; returns (row, index)."""
        # Rates are computed over the whole series, so points near the
        # edges of the selection still get a full centered window.
        segment = rate_values[left:right + 1]
        nearest_local = int(np.argmin(np.abs(segment - target)))
        source_index = left + nearest_local

        row = {'speed': float(segment[nearest_local])}
        for column in self.Current_columns:
            value = pd.to_numeric(
                self.Current.iloc[source_index][column], errors='coerce'
            )
            if pd.isna(value):
                raise ValueError(
                    f'Current column "{column}" contains a non-numeric '
                    f'value at the selected point.'
                )
            # The divisor applies to the found current.
            row[column] = float(value) / divide
            # Pmax = Imax * MC (only if MC is set)
            if mc is not None:
                row[f'Pmax {column}'] = row[column] * mc
        return row, source_index

    def _insert_row(self, row: dict[str, float]) -> None:
        self.tree.insert(
            '', 'end',
            values=[f"{row['speed']:.4f}"] +
                   [f"{row[c]:.6g}" for c in self.Current_columns] +
                   [f"{row[c]:.6g}" for c in self.Pmax_columns]
        )

    def proceed(self) -> None:
        if self.selection_left is None or self.selection_right is None:
            Messages.show('warning', 'LINE_NOT_DEFINED')
            return

        try:
            target_speed_h, divide = self._read_parameters()

            left, right = sorted((self.selection_left, self.selection_right))
            if right - left < 1:
                raise ValueError(
                    'Select an interval containing at least two data points.'
                )

            row, source_index = self._compute_step(
                left, right, target_speed_h, divide, self.rate_values, self.MC
            )
            self.steps.append({
                'left': left, 'right': right, 'target': target_speed_h,
                'divide': divide, 'source_index': source_index, 'row': row,
            })
            self._insert_row(row)
            self.refresh_point_marker()
            self.update_target_line()

            self.selection_label.config(
                text=(
                    f'Target: {target_speed_h:.4f} °C/h; '
                    f'found: {row["speed"]:.4f} °C/h at '
                    f'{self._format_time(self.time_num[source_index])}'
                )
            )
        except Exception as exc:
            Messages.show('error', 'VALUE_ERROR', value_name='Imax', error=exc)

    def remove_previous_step(self) -> None:
        """Delete the last row of the table (and its line on the plot)."""
        if not self.steps:
            Messages.show('warning', 'NO_DRDH_RESULTS')
            return
        self.steps.pop()
        children = self.tree.get_children()
        if children:
            self.tree.delete(children[-1])
        self.refresh_point_marker()

    # ------------------------------------------------------------------
    # Computed parameters (MC, base)
    # ------------------------------------------------------------------
    def show_computed_parameters(self) -> None:
        win = tk.Toplevel(self.root)
        win.title('Computed parameters')
        win.geometry('340x180')
        win.config(bg=COLORS['BACKGROUND_COLOR'])
        frame = tk.Frame(win)
        frame.pack(
            fill=tk.BOTH, expand=True,
            padx=GAPS['GAPS_X']['PAD_X_10'], pady=GAPS['GAPS_Y']['PAD_Y_10']
        )
        listbox = tk.Listbox(
            frame, font=FONTS['DATA_FONT'],
            bg=COLORS['PARAMETERS_BACKGROUND_COLOR']
        )
        listbox.pack(fill=tk.BOTH, expand=True)
        mc_text = 'not set' if self.MC is None else f'{self.MC}'
        listbox.insert(tk.END, f'MC, MWt/(°C/h): {mc_text}')
        listbox.insert(tk.END, f'base, s: {self.base}')
        SmallButtons(win, text='Close', command=win.destroy).pack(pady=5)

    def edit_computed_parameters(self) -> None:
        win = tk.Toplevel(self.root)
        win.title('Edit computed parameters')
        win.geometry('380x170')
        win.config(bg=COLORS['BACKGROUND_COLOR'])
        win.grab_set()

        tk.Label(
            win, text='MC, MWt/(°C/h):', font=FONTS['INFO_FONT'],
            bg=COLORS['BACKGROUND_COLOR']
        ).grid(row=0, column=0, padx=10, pady=(15, 5), sticky='w')
        mc_entry = ttk.Entry(win, width=ENTRY_WIDTH, font=FONTS['DATA_FONT'])
        mc_entry.grid(row=0, column=1, padx=10, pady=(15, 5))
        if self.MC is not None:
            mc_entry.insert(0, str(self.MC))

        tk.Label(
            win, text='base, s:', font=FONTS['INFO_FONT'],
            bg=COLORS['BACKGROUND_COLOR']
        ).grid(row=1, column=0, padx=10, pady=5, sticky='w')
        base_entry = ttk.Entry(win, width=ENTRY_WIDTH, font=FONTS['DATA_FONT'])
        base_entry.grid(row=1, column=1, padx=10, pady=5)
        base_entry.insert(0, str(self.base))
        mc_entry.focus_set()

        def apply() -> None:
            # MC: empty means "no MC" (no Pmax columns)
            mc_text = mc_entry.get().strip().replace(',', '.')
            mc: Optional[float] = None
            if mc_text:
                try:
                    mc = float(mc_text)
                except ValueError as exc:
                    Messages.show('error', 'VALUE_ERROR', value_name='MC', error=exc)
                    return
                if mc <= 0:
                    Messages.show(
                        'error', 'VALUE_POSTIVE', value='MC', sign='positive'
                    )
                    return

            try:
                base = float(base_entry.get().strip().replace(',', '.'))
            except ValueError as exc:
                Messages.show('error', 'VALUE_ERROR', value_name='base', error=exc)
                return
            low, high = IMAX_BASE_RANGE
            if not (low <= base <= high):
                Messages.show(
                    'warning', 'BASE_OUT_OF_RANGE',
                    value='base', low=low, high=high
                )
                return

            try:
                self.apply_computed_parameters(mc, int(round(base)))
            except Exception as exc:
                Messages.show('error', 'VALUE_ERROR', value_name='Imax', error=exc)
                return
            win.destroy()

        SmallButtons(win, text='Apply', command=apply).grid(
            row=2, column=0, padx=10, pady=20
        )
        SmallButtons(win, text='Cancel', command=win.destroy).grid(
            row=2, column=1, padx=10, pady=20
        )

    def apply_computed_parameters(self, mc: Optional[float], base: int) -> None:
        """Set new MC / base and recalculate everything that depends on them.

        The rows already in the table are found again with the same
        interval, requested speed and divisor, so the table never mixes
        results obtained with different base / MC.
        """
        rate_values = self._heating_rate_per_hour(self.temperature_values, base)
        recalculated = [
            self._compute_step(
                s['left'], s['right'], s['target'], s['divide'], rate_values, mc
            )
            for s in self.steps
        ]

        # Everything was calculated successfully: commit.
        self.MC = mc
        self.base = base
        self.rate_values = rate_values
        self.rate_line.set_ydata(rate_values)
        for step, (row, source_index) in zip(self.steps, recalculated):
            step['row'] = row
            step['source_index'] = source_index

        # Pmax columns appear / disappear together with MC
        self.tree.destroy()
        self.create_table()
        for step in self.steps:
            self._insert_row(step['row'])

        self.ax2.relim()
        self.ax2.autoscale_view(scalex=False)
        self.refresh_point_marker()
        self.update_target_line()

    def on_scroll(self, event) -> None:
        """Zoom the X axis around the mouse cursor; Y axis remains unchanged."""
        if event.inaxes not in self.plot_axes or event.xdata is None:
            return
        if self.toolbar.mode != '':
            return

        x_min, x_max = self.ax.get_xlim()
        x_cursor = float(event.xdata)

        if event.button == 'up':
            scale = 1.0 / self._zoom_factor
        elif event.button == 'down':
            scale = self._zoom_factor
        else:
            return

        new_left = x_cursor - (x_cursor - x_min) * scale
        new_right = x_cursor + (x_max - x_cursor) * scale

        data_min = float(np.nanmin(self.time_num))
        data_max = float(np.nanmax(self.time_num))
        new_left = max(new_left, data_min)
        new_right = min(new_right, data_max)

        if new_right <= new_left:
            return

        self.ax.set_xlim(new_left, new_right)
        self.canvas.draw_idle()

    def get_table_data(self) -> list[list[str]]:
        return [list(self.tree.item(item, 'values')) for item in self.tree.get_children()]

    def get_column_headings(self) -> list[str]:
        return (
                ['Temperature rise speed, °C/h']
                + [f'Imax {c}' for c in self.Current_columns]
                + self.Pmax_columns
        )

    def choose_save_format(self) -> None:
        win = tk.Toplevel(self.root)
        win.title('Save Format')
        win.geometry('300x120')
        win.config(bg=COLORS['BACKGROUND_COLOR'])
        win.grab_set()
        tk.Label(
            win, text='Please, choose file format:', font=FONTS['DATA_FONT'],
            bg=COLORS['BACKGROUND_COLOR']
        ).grid(row=0, column=0, columnspan=2, pady=(10, 20))
        SmallButtons(win, text='.txt', command=lambda: (win.destroy(), self.save_to_txt())).grid(
            row=2, column=0, padx=10
        )
        SmallButtons(win, text='Excel', command=lambda: (win.destroy(), self.save_to_excel())).grid(
            row=2, column=1, padx=10
        )

    def save_to_excel(self) -> None:
        """
        Export the results table to an Excel file.

        Creates a formatted workbook with auto-sized columns.
        """
        if not self.results:
            Messages.show('warning', 'NO_DRDH_RESULTS')
            return
        filename = filedialog.asksaveasfilename(
            initialdir=os.path.dirname(os.path.abspath(__file__)),
            defaultextension='.xlsx',
            filetypes=[('Excel files', '*.xlsx')], title='Save results'
        )
        if not filename:
            return
        wb = Workbook()
        ws = wb.active
        ws.title = 'Results'
        headings = self.get_column_headings()
        ws.append(headings)
        for cell in ws[1]:
            cell.font = FONTS['EXCEL_FONT']
        rows = self.get_table_data()
        for row in rows:
            ws.append(row)
        for idx, heading in enumerate(headings, 1):
            max_len = max(len(str(heading)), *(len(str(r[idx - 1])) for r in rows))
            ws.column_dimensions[get_column_letter(idx)].width = max_len + 2
        wb.save(filename)

    def save_to_txt(self) -> None:
        """
        Export the results table as a tab-separated text file.
        """
        if not self.results:
            Messages.show('warning', 'NO_DRDH_RESULTS')
            return
        filename = filedialog.asksaveasfilename(
            initialdir=os.path.dirname(os.path.abspath(__file__)),
            defaultextension='.txt', filetypes=[('Text files', '*.txt')],
            title='Save results'
        )
        if not filename:
            return
        with open(filename, 'w', encoding='utf-8') as f:
            f.write('\t'.join(self.get_column_headings()) + '\n')
            for row in self.get_table_data():
                f.write('\t'.join(map(str, row)) + '\n')

    def back(self, window: tk.Widget) -> None:
        root_window = self.root.winfo_toplevel()
        root_window.config(menu=None)
        if self.menu_bar is not None:
            self.menu_bar.destroy()
            self.menu_bar = None
        window.destroy()
        self.main_app.create_Imax_window()
        root_window.title('Imax')
        root_window.update_idletasks()
